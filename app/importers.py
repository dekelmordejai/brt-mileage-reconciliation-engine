"""Adaptadores de originales; nunca corrigen ni sobrescriben la fuente."""

import csv
import io
import json
import re
import unicodedata
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from functools import lru_cache

KINDS = {"RV", "PASOS", "FLOTA", "CIRCUITOS", "APOYOS", "DESINC", "TARIFAS", "DISTANCIAS"}


def key(value):
    # Unicode normalization for heterogeneous source headers.
    value = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(
        "[^A-Z0-9]+", " ", "".join(c for c in value if not unicodedata.combining(c)).upper()
    ).strip()


def clean(value):
    value = str(value if value is not None else "").strip()
    return "" if value.upper() in {"-", "NAN", "NONE", "N/A", "NAT"} else value


def identifier(value):
    return re.sub(r"\.0$", "", clean(value))


def number(value):
    value = clean(value).replace(" ", "")
    if not value:
        return None
    if "," in value and "." in value:
        raise ValueError("Separadores numéricos ambiguos; normalice explícitamente")
    try:
        result = Decimal(value.replace(",", "."))
        if not result.is_finite():
            raise ValueError("Número no finito")
        return str(result)
    except InvalidOperation as exc:
        raise ValueError(f"Número inválido: {value}") from exc


@lru_cache(maxsize=10000)
def day(value):
    value = clean(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Fecha inválida: {value}")


@lru_cache(maxsize=10000)
def clock(value):
    value = clean(value)
    if not value:
        return None
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            t = datetime.strptime(value, fmt)
            return t.hour * 60 + t.minute + t.second / 60
        except ValueError:
            pass
    raise ValueError(f"Hora inválida: {value}")


def operational_dates(date, departure):
    d = datetime.fromisoformat(date).date()
    t = clock(departure)
    return {
        "conventional": (d - timedelta(days=int(t is not None and t < 200))).isoformat(),
        "provisional": (d - timedelta(days=int(t is not None and t < 180))).isoformat(),
        "boundary_pending": t is not None and 180 <= t < 200,
    }


def read_table(data, filename, sheet=None):
    if filename.lower().endswith(".json"):
        rows = json.loads(data)
        if not isinstance(rows, list):
            raise ValueError("JSON debe contener una lista de filas")
        return rows
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = book[sheet] if sheet else book.worksheets[0]
        values = ws.iter_rows(values_only=True)
        headers = [clean(v) for v in next(values)]
        return [
            dict(zip(headers, [clean(v) for v in row], strict=False))
            for row in values
            if any(v is not None for v in row)
        ]
    if not filename.lower().endswith(".csv"):
        raise ValueError(
            "Use CSV, XLSX o JSON; PDF requiere transcripción versionada, no inferencia automática"
        )
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Codificación no compatible")
    delimiter = ";" if text.splitlines()[0].count(";") >= text.splitlines()[0].count(",") else ","
    return list(csv.DictReader(io.StringIO(text), delimiter=delimiter))


def normalize(kind, raw):
    r = {key(k): clean(v) for k, v in raw.items()}

    def get(*names):
        return next((r[key(n)] for n in names if r.get(key(n))), "")

    def required(*names):
        v = get(*names)
        if not v:
            raise ValueError("Falta " + names[0])
        return v

    if kind == "FLOTA":
        return dict(
            vehicle=identifier(required("vehicle", "económico")),
            company=required("company", "EO", "empresa"),
            vehicle_type=required("vehicle_type", "tipo de autobús", "tipo", "autobus"),
            valid_from=day(get("valid_from") or "1900-01-01"),
            valid_to=day(get("valid_to") or "9999-12-31"),
        )
    if kind == "TARIFAS":
        rate = number(required("rate", "tarifa"))
        if Decimal(rate) <= 0:
            raise ValueError("Tarifa debe ser positiva")
        return dict(
            company=required("company", "empresa"),
            vehicle_type=required("vehicle_type", "tipo"),
            line=required("line", "línea"),
            valid_from=day(required("valid_from", "desde")),
            valid_to=day(required("valid_to", "hasta")),
            rate=rate,
            reference=required("reference", "autorización"),
        )
    if kind == "DISTANCIAS":
        return dict(
            route=required("route", "ruta"),
            start=required("start", "origen"),
            end=required("end", "destino"),
            km=number(required("km", "km_tramo")),
        )
    if kind in {"RV", "CIRCUITOS"}:
        date = day(required("date", "fecha"))
        route = required("route", "ruta", "circuito").split(" ")[0]
        dep = get("departure", "partida real")
        planned = get("planned_departure", "partida planeada")
        clock(dep)
        clock(planned)
        arrival = get("arrival", "llegada real")
        clock(arrival)
        return dict(
            date=date,
            route=route,
            line="L" + str(int(route[1:3])) if re.match(r"[LC]\d{2}", route) else get("line"),
            journey=identifier(get("journey", "jornada")),
            trip=identifier(get("trip", "viaje")),
            planned_company=get("planned_company", "empresa"),
            vehicle=identifier(get("vehicle", "vehiculo real", "económico")),
            assigned_vehicle=identifier(get("assigned_vehicle", "económico planificado")),
            planned_type=get("planned_type"),
            departure=dep,
            planned_departure=planned,
            arrival=arrival,
            planned_km=number(get("planned_km", "KM Planeada")),
            physical_km=number(get("physical_km", "KM Realizado")),
            status=get("status", "desc. status del viaje"),
            dates=operational_dates(date, planned or dep),
        )
    if kind == "PASOS":
        checkpoints = {
            k: v
            for k, v in raw.items()
            if key(k)
            not in {
                "FECHA",
                "DATE",
                "JORNADA",
                "JOURNEY",
                "ECONOMICO",
                "VEHICLE",
                "VIAJE",
                "TRIP",
                "KM REALIZADO",
                "PHYSICAL KM",
            }
            and clean(v)
        }
        for v in checkpoints.values():
            clock(v)
        return dict(
            date=day(required("date", "fecha")),
            journey=identifier(get("journey", "jornada")),
            vehicle=identifier(required("vehicle", "económico")),
            trip=get("trip", "viaje"),
            physical_km=number(get("physical_km", "km realizado")),
            checkpoints=checkpoints,
            checkpoint_order=list(checkpoints),
        )
    date = day(
        required("date", "fecha de apoyo" if kind == "APOYOS" else "fecha de desincorporación")
    )
    time = required("time", "hora de apoyo" if kind == "APOYOS" else "hora de desincorporación")
    clock(time)
    end = get(
        "end_time", "hora de reincorporación" if kind == "DESINC" else "hora de desincorporación"
    )
    clock(end)
    end_date = get(
        "end_date", "fecha de reincorporación" if kind == "DESINC" else "fecha de desincorporación"
    )
    km = number(get("km2" if kind == "APOYOS" else "km"))
    return dict(
        date=date,
        time=time,
        end_time=end,
        end_date=day(end_date) if end_date else "",
        line=get("line", "línea"),
        journey=identifier(get("journey", "jornada")),
        vehicle=identifier(required("vehicle", "económico")),
        replacement=get("replacement", "económico de reincorporación"),
        company=get("company", "EO"),
        station=get("station", "estación de apoyo", "estación de desincorporación"),
        km=None if kind == "APOYOS" and km is not None and Decimal(km) == 0 else km,
        reason=get("reason", "motivo"),
        verified=get("verified", "verificado"),
    )
