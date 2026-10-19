import io
import json
import sqlite3
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.importers import normalize, operational_dates
from app.main import create_app
from app.security import password_hash


@pytest.fixture
def system(tmp_path):
    app = create_app(tmp_path / "db.sqlite3")
    with app.state.store.connect() as db:
        for name, role in [
            ("admin", "admin"),
            ("analyst", "analyst"),
            ("reviewer", "reviewer"),
            ("viewer", "viewer"),
        ]:
            db.execute(
                "INSERT INTO users(id,password,role) VALUES(?,?,?)",
                (name, password_hash("Test-local-2026!"), role),
            )
    client = TestClient(app)
    login(client, "admin")
    return client, app


def login(client, user):
    r = client.post("/api/login", json={"username": user, "password": "Test-local-2026!"})
    assert r.status_code == 200, r.text
    client.headers["X-CSRF-Token"] = r.json()["csrf"]


def seed(client, change=None):
    data = json.loads((Path(__file__).parents[1] / "fixtures" / "synthetic.json").read_text())
    if change:
        change(data)
    for kind, rows in data.items():
        r = client.post(
            "/api/sources",
            data={"kind": kind},
            files={"file": (kind + ".json", json.dumps(rows).encode())},
        )
        assert r.status_code == 200, r.text
        assert r.json().get("invalid", 0) == 0, r.text
    return [s["id"] for s in client.get("/api/sources").json()]


def run(client, sources):
    r = client.post("/api/runs", json={"line": "L1", "year": 2026, "week": 40, "sources": sources})
    assert r.status_code == 201, r.text
    rid = r.json()["id"]
    result = client.get("/api/runs/" + rid).json()
    rows = client.get("/api/runs/" + rid + "/rows").json()["rows"]
    return rid, result, rows


def proposal(client, row, **changes):
    payload = {
        "rv_id": row["id"],
        "parent": row["approved_proposal"],
        "segments": [
            {
                "start_km": "0",
                "end_km": row["physical_km"],
                "origin": "Estación A",
                "destination": "Estación B",
                "evidence_ids": row["candidates"],
            }
        ],
        "reason": "Validación documental sintética",
        "evidence": "Expediente sintético QA-001",
    }
    payload.update(changes)
    return client.post("/api/proposals", json=payload)


def approve(client, pid):
    login(client, "reviewer")
    r = client.post(
        f"/api/proposals/{pid}/decision",
        json={"approved": True, "reason": "Evidencia y alcance revisados por segunda persona"},
    )
    assert r.status_code == 200, r.text
    login(client, "analyst")


def test_full_import_review_rerun_exports_and_persistence(system):
    c, app = system
    sources = seed(c)
    rid, r, rows = run(c, sources)
    assert r["totals"]["planned_km"] == "30"
    assert r["totals"]["physical_km"] == "30"  # Circuito no suma otra vez.
    assert r["totals"]["eligible_km"] == "0"
    assert "CANCELADA_CON_KM_FISICO" in rows[0]["alerts"]
    assert rows[0]["assigned_vehicle"] == "DEMO-999"
    login(c, "analyst")
    p = proposal(c, rows[0])
    assert p.status_code == 201, p.text
    login(c, "admin")
    # Una persona distinta puede revisar, el autor no puede autoaprobar.
    login(c, "analyst")
    assert (
        c.post(
            "/api/proposals/" + p.json()["id"] + "/decision",
            json={"approved": True, "reason": "Intento de autoaprobación"},
        ).status_code
        == 403
    )
    approve(c, p.json()["id"])
    rid2, r2, rows2 = run(c, sources)
    assert r2["totals"]["physical_km"] == "30"
    assert r2["totals"]["eligible_km"] == "10"
    assert r2["totals"]["estimated_amount"] == "102.50"
    assert (
        c.get("/api/runs/" + rid).json()["totals"]["eligible_km"] == "0"
    )  # Snapshot anterior intacto.
    assert c.get("/api/runs/" + rid2 + "/export/csv").status_code == 200
    x = c.get("/api/runs/" + rid2 + "/export/xlsx")
    book = load_workbook(io.BytesIO(x.content))
    assert {"Viajes", "Desgloses", "Resumen", "Advertencias", "Fuentes"} <= set(book.sheetnames)
    login(c, "reviewer")
    assert c.get("/api/audit").json()["chain_valid"]
    other = create_app(app.state.store.path)
    with other.state.store.connect() as db:
        assert db.execute("SELECT count(*) FROM runs").fetchone()[0] == 2
        for table in ("sources", "records", "runs", "proposals", "decisions", "audit"):
            with pytest.raises(sqlite3.IntegrityError):
                db.execute(f"DELETE FROM {table}")


def test_tariff_real_type_and_local_event_regularization(system):
    c, _ = system
    sources = seed(c)
    _, _, rows = run(c, sources)
    transfer = rows[1]
    login(c, "analyst")
    p = proposal(c, transfer)
    assert p.status_code == 201
    approve(c, p.json()["id"])
    _, r, rs = run(c, sources)
    assert rs[1]["rate"] == "15.50"
    assert rs[1]["eligible_km"] == "0"  # Apoyo pendiente localizado.
    assert "EVENTO_CCO_LOCALIZADO_PENDIENTE" not in rs[0]["blocks"]
    correction = proposal(c, rs[1], cleared_events=rs[1]["events"])
    assert correction.status_code == 201, correction.text
    approve(c, correction.json()["id"])
    _, r, rs = run(c, sources)
    assert rs[1]["eligible_km"] == "5"
    assert rs[1]["estimated_amount"] == "77.50"
    assert "TRASLADO_PROGRAMADO_VALIDABLE" in rs[1]["alerts"]
    assert r["totals"]["physical_km"] == "30"
    with c:
        pass


def test_cross_company_not_automatically_credited(system):
    c, _ = system
    s = seed(c)
    _, _, rows = run(c, s)
    login(c, "analyst")
    p = proposal(c, rows[3])
    approve(c, p.json()["id"])
    _, r, rows = run(c, s)
    assert rows[3]["covered_km"] == "9"
    assert rows[3]["eligible_km"] == "0"
    assert "INTEREMPRESA_ATRIBUCION_PENDIENTE" in rows[3]["blocks"]


def test_exclusive_evidence_segments_and_stale_proposals(system):
    c, _ = system
    s = seed(c)
    _, _, rows = run(c, s)
    login(c, "analyst")
    seg = {
        "start_km": "0",
        "end_km": "6",
        "origin": "A",
        "destination": "B",
        "evidence_ids": rows[0]["candidates"],
    }
    assert proposal(c, rows[0], segments=[seg, seg]).status_code == 422
    p = proposal(c, rows[0])
    second = proposal(c, rows[0])
    approve(c, p.json()["id"])
    login(c, "reviewer")
    assert (
        c.post(
            "/api/proposals/" + second.json()["id"] + "/decision",
            json={"approved": True, "reason": "Versión desactualizada detectada"},
        ).status_code
        == 409
    )
    login(c, "analyst")
    assert proposal(c, rows[2], segments=[seg]).status_code == 409


def test_deduplication_and_overlapping_vehicle_block(system):
    c, _ = system

    def change(d):
        d["RV"].append(dict(d["RV"][0]))
        d["RV"].append(dict(d["RV"][0], trip="conflict", departure="06:10", arrival="06:25"))

    s = seed(c, change)
    _, r, rows = run(c, s)
    assert r["totals"]["physical_km"] == "40"  # Duplicado no se suma, conflicto se conserva físico.
    assert "DUPLICADO_RV" in rows[0]["blocks"]
    assert "SOLAPAMIENTO_VEHICULO" in rows[0]["blocks"]
    assert r["totals"]["eligible_km"] == "0"


def test_schema_quarantine_idempotency_and_csrf(system):
    c, _ = system
    data = json.dumps([{"date": "incorrecta", "route": "L01A01-1"}]).encode()
    r = c.post("/api/sources", data={"kind": "RV"}, files={"file": ("bad.json", data)})
    assert r.json()["invalid"] == 1
    assert c.post("/api/sources", data={"kind": "RV"}, files={"file": ("bad.json", data)}).json()[
        "duplicate"
    ]
    assert c.post("/api/runs", json={"sources": [r.json()["id"]]}).status_code == 422
    del c.headers["X-CSRF-Token"]
    assert c.post("/api/runs", json={"sources": ["x"]}).status_code == 403
    assert c.get("/api/me", headers={"host": "evil.example"}).status_code == 400
    login(c, "viewer")
    assert c.post("/api/runs", json={"sources": ["x"]}).status_code == 403
    login(c, "analyst")
    assert c.get("/api/users").status_code == 403


def test_semantics_and_midnight_boundaries():
    row = normalize("APOYOS", {"date": "2026-09-28", "time": "23:00", "vehicle": "SYN", "km2": "0"})
    assert row["km"] is None
    assert operational_dates("2026-09-29", "03:10") == {
        "conventional": "2026-09-28",
        "provisional": "2026-09-29",
        "boundary_pending": True,
    }
    assert operational_dates("2026-09-29", "03:20")["conventional"] == "2026-09-29"
    assert operational_dates("2026-09-29", "02:59")["provisional"] == "2026-09-28"


def test_tariff_overlap_preserves_physical_and_blocks_economics(system):
    c, _ = system

    def change(d):
        d["TARIFAS"].append(dict(d["TARIFAS"][0], reference="SYN-OVERLAP"))

    s = seed(c, change)
    _, r, rows = run(c, s)
    assert "TARIFA_AUSENTE_O_VIGENCIA_AMBIGUA" in rows[0]["blocks"]
    assert r["totals"]["physical_km"] == "30"


def test_localized_hold_keeps_unaffected_segments_eligible(system):
    c, _ = system
    sources = seed(c)
    _, _, rows = run(c, sources)
    circuit = rows[2]
    login(c, "analyst")
    segments = []
    for a, b, t1, t2 in [
        ("0", "2", "08:00", "08:04"),
        ("2", "4", "08:05", "08:10"),
        ("4", "6", "08:11", "08:20"),
    ]:
        segments.append(
            {
                "start_km": a,
                "end_km": b,
                "origin": "Estación A",
                "destination": "Estación B",
                "departure": t1,
                "arrival": t2,
                "evidence_ids": circuit["candidates"],
            }
        )
    p = proposal(c, circuit, segments=segments)
    assert p.status_code == 201, p.text
    approve(c, p.json()["id"])
    _, r, rows = run(c, sources)
    assert rows[2]["physical_km"] == "6"
    assert rows[2]["covered_km"] == "6"
    assert rows[2]["eligible_km"] == "4"
    assert [s["eligible_km"] for s in rows[2]["segment_ledger"]] == ["2", "0", "2"]
    assert r["totals"]["physical_km"] == "30"


def test_approved_evidence_must_remain_in_selected_sources(system):
    c, _ = system
    sources = seed(c)
    _, _, rows = run(c, sources)
    login(c, "analyst")
    p = proposal(c, rows[0])
    approve(c, p.json()["id"])
    login(c, "admin")
    replacement = [
        {
            "date": "2026-09-28",
            "journey": "SYN-OTHER",
            "vehicle": "SYN-OTHER",
            "A": "12:00",
            "B": "12:30",
        }
    ]
    imported = c.post(
        "/api/sources",
        data={"kind": "PASOS"},
        files={"file": ("other.json", json.dumps(replacement).encode())},
    ).json()
    old_pasos = {
        s["id"]
        for s in c.get("/api/sources").json()
        if s["kind"] == "PASOS" and s["id"] != imported["id"]
    }
    new_sources = [s for s in sources if s not in old_pasos] + [imported["id"]]
    _, r, rows = run(c, new_sources)
    assert rows[0]["eligible_km"] == "0"
    assert "EVIDENCIA_AUTORIZADA_FUERA_DE_FUENTES" in rows[0]["blocks"]
